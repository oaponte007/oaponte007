"""Builds the branded .docx report. Needs `python-docx` -- unlike the
collectors (which must run dependency-free on the target itself), this
generation step is meant to run wherever's convenient: the same
machine if it has python-docx installed, or an admin's own workstation
after pulling the collected JSON off the target. See README.

Deliberately NOT reusing chpc_helper_core's engine/wizard (this isn't a
choice-driven template generator the way the other three chpc-*-helper
apps are -- it's a fixed report layout filled with real findings), so
the only shared code with the rest of this repo is the Coastal HPC
branding constants in branding.py, itself already established in
billing-statement-builder/.
"""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from . import branding
from .report_data import ReportFinding, sort_findings, summarize_by_severity


def _set_cell_shading(cell, hex_color: str) -> None:
    shd = cell._tc.get_or_add_tcPr().makeelement(qn("w:shd"), {
        qn("w:val"): "clear", qn("w:color"): "auto", qn("w:fill"): hex_color,
    })
    cell._tc.get_or_add_tcPr().append(shd)


def _add_centered_picture(doc: Document, image_path: Path, width_inches: float) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    run.add_picture(str(image_path), width=Inches(width_inches))


def _add_severity_summary_table(doc: Document, counts: dict[str, int]) -> None:
    present = [s for s in branding.SEVERITY_ORDER if counts.get(s)]
    if not present:
        return
    table = doc.add_table(rows=1, cols=len(present))
    table.style = "Table Grid"
    header = table.rows[0].cells
    for i, severity in enumerate(present):
        header[i].text = f"{branding.SEVERITY_LABELS[severity]}\n{counts[severity]}"
        _set_cell_shading(header[i], branding.SEVERITY_COLORS[severity])
        for p in header[i].paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.color.rgb = RGBColor.from_string("FFFFFF")
                run.font.bold = True


def _add_finding(doc: Document, finding: ReportFinding, index: int) -> None:
    heading = doc.add_paragraph()
    badge_run = heading.add_run(f"[{branding.SEVERITY_LABELS[finding.severity]}] ")
    badge_run.font.bold = True
    badge_run.font.color.rgb = RGBColor.from_string(branding.SEVERITY_COLORS[finding.severity])
    title_run = heading.add_run(finding.title)
    title_run.font.bold = True
    title_run.font.size = Pt(13)

    if finding.cve_ids:
        cve_p = doc.add_paragraph()
        cve_run = cve_p.add_run(f"Reference: {finding.advisory_id or ''}  ({', '.join(finding.cve_ids)})")
        cve_run.font.italic = True
        cve_run.font.size = Pt(9)

    doc.add_paragraph(finding.description)

    danger_p = doc.add_paragraph()
    danger_label = danger_p.add_run("Why this matters: ")
    danger_label.font.bold = True
    danger_p.add_run(finding.danger)

    fix_label_p = doc.add_paragraph()
    fix_label_run = fix_label_p.add_run("How to fix it:")
    fix_label_run.font.bold = True
    for step in finding.fix_steps:
        doc.add_paragraph(step, style="List Number")

    if finding.evidence:
        evidence_p = doc.add_paragraph()
        evidence_run = evidence_p.add_run(f"What was found: {finding.evidence}")
        evidence_run.font.italic = True
        evidence_run.font.size = Pt(9)
        evidence_run.font.color.rgb = RGBColor.from_string("666666")

    doc.add_paragraph()  # spacing before the next finding


def build_docx_report(
    hostname: str,
    os_label: str,
    scan_date: str,
    findings: list[ReportFinding],
    feed_notes: list[str],
    output_path: Path,
) -> Path:
    doc = Document()

    logo_path = branding.ensure_logo()
    if logo_path:
        _add_centered_picture(doc, logo_path, width_inches=2.2)

    title = doc.add_heading("Security Baseline & Vulnerability Report", level=1)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    slogan_p = doc.add_paragraph()
    slogan_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    slogan_run = slogan_p.add_run(f"{branding.COMPANY_NAME} — {branding.SLOGAN}")
    slogan_run.font.italic = True
    slogan_run.font.color.rgb = RGBColor.from_string(branding.PRIMARY_HEX)

    website_p = doc.add_paragraph()
    website_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    website_run = website_p.add_run(branding.WEBSITE)
    website_run.font.size = Pt(9)
    website_run.font.color.rgb = RGBColor.from_string(branding.ACCENT_HEX)

    doc.add_paragraph()

    info_table = doc.add_table(rows=3, cols=2)
    info_table.style = "Table Grid"
    rows = [("Host", hostname), ("Operating System", os_label), ("Report generated", scan_date)]
    for i, (label, value) in enumerate(rows):
        info_table.rows[i].cells[0].text = label
        info_table.rows[i].cells[1].text = value
        info_table.rows[i].cells[0].paragraphs[0].runs[0].font.bold = True

    doc.add_paragraph()
    doc.add_heading("Executive Summary", level=2)

    sorted_findings = sort_findings(findings)
    counts = summarize_by_severity(sorted_findings)
    total = len(sorted_findings)

    if total == 0:
        doc.add_paragraph(
            "No baseline hardening issues or known vulnerabilities were found against the data "
            "available at scan time. Keep this computer's updates and settings current -- this is "
            "a snapshot, not a permanent guarantee."
        )
    else:
        doc.add_paragraph(f"{total} finding(s) were identified on this host, summarized by severity below.")
        _add_severity_summary_table(doc, counts)

    if feed_notes:
        doc.add_paragraph()
        notes_heading = doc.add_paragraph()
        notes_heading.add_run("Notes on data currency:").font.bold = True
        for note in feed_notes:
            doc.add_paragraph(note, style="List Bullet")

    if sorted_findings:
        doc.add_paragraph()
        doc.add_heading("Detailed Findings", level=2)
        doc.add_paragraph(
            "Findings are listed from most to least urgent. Each one explains what it is, why it "
            "matters, and exactly how to fix it -- written so anyone can follow the steps, not just "
            "an IT professional. Fixing these is left up to you and your team; nothing on this "
            "computer was changed by running this scan."
        )
        current_severity = None
        for i, finding in enumerate(sorted_findings, 1):
            if finding.severity != current_severity:
                current_severity = finding.severity
                doc.add_heading(f"{branding.SEVERITY_LABELS[current_severity]} severity", level=3)
            _add_finding(doc, finding, i)

    doc.add_paragraph()
    footer_p = doc.add_paragraph()
    footer_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer_run = footer_p.add_run(
        f"{branding.COMPANY_NAME} — {branding.SLOGAN}  |  {branding.WEBSITE}"
    )
    footer_run.font.size = Pt(9)
    footer_run.font.color.rgb = RGBColor.from_string("888888")

    output_path = Path(output_path)
    doc.save(str(output_path))
    return output_path
