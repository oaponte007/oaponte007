from pathlib import Path

import docx as docx_lib

from chpc_sec_vuln_helper.feeds.matcher import match_linux
from chpc_sec_vuln_helper.feeds.normalize_rhel_oval import parse_oval
from chpc_sec_vuln_helper.report import report_data
from chpc_sec_vuln_helper.rules.engine import evaluate
from chpc_sec_vuln_helper.rules.linux_rules import LINUX_RULES

FIXTURE = Path(__file__).parent / "fixtures" / "sample_oval.xml"


def _sample_findings():
    bad = {
        "os": {"family": "rhel", "major_version": 9},
        "selinux_status": "disabled",
        "firewall_active": False,
    }
    baseline = evaluate(LINUX_RULES, bad)
    vulns = parse_oval(FIXTURE.read_text())
    installed = [
        {"name": "openssl", "epoch": "1", "version": "3.0.7", "release": "18.el9_4", "arch": "x86_64"},
        {"name": "openssl-libs", "epoch": "1", "version": "3.0.7", "release": "18.el9_4", "arch": "x86_64"},
    ]
    matches = match_linux(installed, vulns)
    return report_data.from_baseline_findings(baseline) + report_data.from_cve_matches(matches, "rhel")


def test_cve_matches_for_same_advisory_are_grouped_into_one_finding():
    vulns = parse_oval(FIXTURE.read_text())
    installed = [
        {"name": "openssl", "epoch": "1", "version": "3.0.7", "release": "18.el9_4", "arch": "x86_64"},
        {"name": "openssl-libs", "epoch": "1", "version": "3.0.7", "release": "18.el9_4", "arch": "x86_64"},
    ]
    matches = match_linux(installed, vulns)
    assert len(matches) == 2  # matcher still reports both package matches...

    findings = report_data.from_cve_matches(matches, "rhel")
    assert len(findings) == 1  # ...but the report groups them into one finding
    assert "openssl" in findings[0].evidence
    assert "openssl-libs" in findings[0].evidence


def test_cve_fix_command_lists_all_affected_packages():
    vulns = parse_oval(FIXTURE.read_text())
    installed = [
        {"name": "openssl", "epoch": "1", "version": "3.0.7", "release": "18.el9_4", "arch": "x86_64"},
        {"name": "openssl-libs", "epoch": "1", "version": "3.0.7", "release": "18.el9_4", "arch": "x86_64"},
    ]
    matches = match_linux(installed, vulns)
    findings = report_data.from_cve_matches(matches, "rhel")
    fix_text = " ".join(findings[0].fix_steps)
    assert "openssl" in fix_text and "openssl-libs" in fix_text


def test_sort_findings_orders_critical_first():
    findings = _sample_findings()
    sorted_findings = report_data.sort_findings(findings)
    severities = [f.severity for f in sorted_findings]
    ranks = {"critical": 0, "high": 1, "medium": 2, "low": 3, "unknown": 4}
    assert severities == sorted(severities, key=lambda s: ranks[s])


def test_summarize_by_severity_counts_match():
    findings = _sample_findings()
    counts = report_data.summarize_by_severity(findings)
    assert sum(counts.values()) == len(findings)


def _full_text(doc) -> str:
    """doc.paragraphs alone skips table cell content (where the
    host/OS info table lives) -- walk both.
    """
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                parts.append(cell.text)
    return "\n".join(parts)


def test_docx_report_builds_and_contains_expected_content(tmp_path):
    findings = _sample_findings()
    from chpc_sec_vuln_helper.report.docx_report import build_docx_report
    out_path = build_docx_report(
        "node042", "Red Hat Enterprise Linux 9.4", "2026-09-29 12:00 UTC",
        findings, ["CVE data: 2 advisories, fetched just now."], tmp_path / "report.docx",
    )
    assert out_path.exists()

    doc = docx_lib.Document(str(out_path))
    full_text = _full_text(doc)
    assert "node042" in full_text
    assert "Coastal HPC" in full_text
    assert "Where High Performance Meets High Security." in full_text
    assert "www.coastal-hpc.com" in full_text
    assert "SELinux" in full_text
    assert "openssl" in full_text
    # the logo image should be embedded
    assert len(doc.inline_shapes) >= 1


def test_docx_report_with_zero_findings_has_reassuring_message(tmp_path):
    from chpc_sec_vuln_helper.report.docx_report import build_docx_report
    out_path = build_docx_report(
        "cleanhost", "Rocky Linux 9.4", "2026-09-29 12:00 UTC", [], [], tmp_path / "report.docx",
    )
    doc = docx_lib.Document(str(out_path))
    full_text = "\n".join(p.text for p in doc.paragraphs)
    assert "No baseline hardening issues or known vulnerabilities" in full_text


def test_html_report_builds_and_contains_expected_content(tmp_path):
    findings = _sample_findings()
    from chpc_sec_vuln_helper.report.html_report import build_html_report
    out_path = build_html_report(
        "node042", "Red Hat Enterprise Linux 9.4", "2026-09-29 12:00 UTC",
        findings, ["CVE data: 2 advisories, fetched just now."], tmp_path / "report.html",
    )
    text = out_path.read_text()
    assert "node042" in text
    assert "Coastal HPC" in text
    assert "www.coastal-hpc.com" in text
    assert "data:image/png;base64," in text  # logo embedded, self-contained
    assert "<html" in text.lower()
    # every severity present must actually render its badge color
    for f in findings:
        from chpc_sec_vuln_helper.report import branding
        assert branding.SEVERITY_COLORS[f.severity] in text


def test_html_report_escapes_html_special_characters(tmp_path):
    from chpc_sec_vuln_helper.report.html_report import build_html_report
    from chpc_sec_vuln_helper.report.report_data import ReportFinding
    tricky = ReportFinding(
        kind="baseline", severity="high", title="<script>alert(1)</script>",
        description="a & b < c", danger="danger & stuff", fix_steps=["step <1>"],
    )
    out_path = build_html_report("h", "os", "date", [tricky], [], tmp_path / "r.html")
    text = out_path.read_text()
    assert "<script>alert(1)</script>" not in text
    assert "&lt;script&gt;" in text
