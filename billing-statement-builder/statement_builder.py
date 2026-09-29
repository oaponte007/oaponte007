"""Fills the Coastal HPC billing statement template with real data.

Kept separate from the GUI (billing_app.py) so the document-generation
logic can be tested/used on its own -- the GUI just collects a dict and
calls build_statement().

The template's placeholder cells hold either:
  (a) a single self-contained value (e.g. a table cell that is just
      "$0.00" or "[MM/DD]") -- replaced wholesale via set_cell_text(), or
  (b) a mix of literal label text and a bracketed placeholder sharing one
      run (e.g. "Statement #: [000001]") -- replaced via a targeted
      substring swap on just the run(s) containing that bracket text, so
      the surrounding label/formatting is untouched.
Word stores a Shift+Enter line break as <w:br/> inside a run; python-docx's
Run.text getter/setter round-trip that as a literal "\\n", which is why
plain string replace on run.text is safe here without special-casing line
breaks.
"""

from __future__ import annotations

import io
import os
import shutil
import tempfile
import zipfile
from copy import deepcopy
from dataclasses import dataclass, field
from datetime import date
from pathlib import Path
from typing import List, Optional, Tuple

import docx
from PIL import Image

LOGO_MEDIA_PATH = "word/media/image1.png"  # the only image in the template: the logo


@dataclass
class ActivityLine:
    date: str = ""
    reference: str = ""          # invoice # / "PAYMENT" / etc.
    description: str = ""
    charge: Optional[float] = None
    payment: Optional[float] = None


@dataclass
class StatementData:
    # statement header
    statement_number: str = ""
    statement_date: str = ""

    # your company (Coastal HPC)
    company_name: str = "COASTAL HPC"
    company_address: str = ""
    company_phone: str = ""
    company_email: str = ""
    company_website: str = ""

    # bill to
    customer_name: str = ""
    contact_name: str = ""
    billing_address: str = ""
    customer_city_state_zip: str = ""
    customer_email: str = ""
    customer_phone: str = ""

    # account summary. previous_balance and past_due are the only figures
    # entered by hand (they come from your books, not from this form).
    # payments_credits/new_charges/amount_due are ALWAYS derived from the
    # activity list by compute_totals() below -- never set them directly,
    # or they can silently disagree with what the activity table shows
    # (this is exactly the bug reported against the first draft: a
    # separately-typed Payments/Credits number that didn't reflect a
    # payment entered as an activity line).
    previous_balance: float = 0.0
    past_due: float = 0.0
    payments_credits: float = 0.0
    new_charges: float = 0.0
    amount_due: float = 0.0

    activity: List[ActivityLine] = field(default_factory=list)

    # Path to a company logo image (PNG/JPG/etc.) to swap into the
    # template in place of its default logo, at the exact same on-page
    # size. Leave None/empty to keep the template's own logo untouched.
    company_logo_path: Optional[str] = None

    payment_methods: str = "ACH / Check / Card / Online"
    payable_to: str = "Coastal HPC"
    payment_instructions: str = ""
    payment_due_date: str = ""

    notes_terms: str = ""


def money(value: float) -> str:
    sign = "-" if value < 0 else ""
    return f"{sign}${abs(value):,.2f}"


def compute_totals(data: StatementData) -> None:
    """Derives payments_credits/new_charges/amount_due from the activity
    lines + the entered previous balance. This is the single source of
    truth for the math -- called live by the GUI on every edit and again
    just before writing, so the displayed numbers and the generated
    document can never disagree, and a payment entered as an activity
    line always reduces the amount due (correctly going negative -- a
    credit balance -- if payments exceed the balance plus new charges)."""
    data.new_charges = sum(a.charge or 0.0 for a in data.activity)
    data.payments_credits = sum(a.payment or 0.0 for a in data.activity)
    data.amount_due = data.previous_balance - data.payments_credits + data.new_charges


# ---------------------------------------------------------------------------
# low-level docx helpers
# ---------------------------------------------------------------------------

def set_paragraph_text(paragraph, text: str) -> None:
    if paragraph.runs:
        paragraph.runs[0].text = text
        for r in paragraph.runs[1:]:
            r.text = ""
    else:
        paragraph.add_run(text)


def set_cell_text(cell, text: str) -> None:
    paragraphs = cell.paragraphs
    set_paragraph_text(paragraphs[0], text)
    for extra_p in paragraphs[1:]:
        for r in extra_p.runs:
            r.text = ""


def replace_in_cell(cell, replacements: dict) -> None:
    """Substring-replaces each old->new pair in every run of every
    paragraph in `cell`, leaving runs that don't contain a match (and thus
    their formatting/line breaks) completely untouched."""
    for p in cell.paragraphs:
        for r in p.runs:
            text = r.text
            changed = False
            for old, new in replacements.items():
                if old in text:
                    text = text.replace(old, new)
                    changed = True
            if changed:
                r.text = text


def replace_in_paragraph(paragraph, replacements: dict) -> None:
    for r in paragraph.runs:
        text = r.text
        changed = False
        for old, new in replacements.items():
            if old in text:
                text = text.replace(old, new)
                changed = True
        if changed:
            r.text = text


def clone_row_before(table, template_row_index: int, before_row_index: int):
    src_tr = table.rows[template_row_index]._tr
    new_tr = deepcopy(src_tr)
    table.rows[before_row_index]._tr.addprevious(new_tr)


def remove_row(table, row_index: int) -> None:
    tr = table.rows[row_index]._tr
    tr.getparent().remove(tr)


# ---------------------------------------------------------------------------
# company logo swap
# ---------------------------------------------------------------------------

def _read_media_pixel_size(docx_path: str, media_path: str) -> Tuple[int, int]:
    with zipfile.ZipFile(docx_path, "r") as z:
        raw = z.read(media_path)
    with Image.open(io.BytesIO(raw)) as img:
        return img.size


def compose_logo_png(logo_source_path: str, canvas_size: Tuple[int, int]) -> bytes:
    """Fits `logo_source_path` onto a transparent canvas of exactly
    `canvas_size` pixels, preserving the logo's own aspect ratio (never
    stretched/distorted) and centering it. `canvas_size` is the template's
    *existing* logo image's own pixel size, whose aspect ratio already
    matches the fixed on-page box (the <wp:extent> in the drawing XML) --
    so any logo composed onto a same-sized canvas fills that exact same
    box when it replaces the original image file."""
    canvas_w, canvas_h = canvas_size
    with Image.open(logo_source_path) as src:
        logo = src.convert("RGBA")
        scale = min(canvas_w / logo.width, canvas_h / logo.height)
        new_w = max(1, round(logo.width * scale))
        new_h = max(1, round(logo.height * scale))
        logo = logo.resize((new_w, new_h), Image.LANCZOS)

    canvas = Image.new("RGBA", (canvas_w, canvas_h), (255, 255, 255, 0))
    offset = ((canvas_w - new_w) // 2, (canvas_h - new_h) // 2)
    canvas.paste(logo, offset, logo)
    buf = io.BytesIO()
    canvas.save(buf, format="PNG")
    return buf.getvalue()


def replace_logo(docx_path: str, logo_source_path: str, media_path: str = LOGO_MEDIA_PATH) -> None:
    """Swaps the template's embedded logo image for `logo_source_path`,
    composed onto a canvas matching the original logo's own pixel size
    (see compose_logo_png) -- so it fills the exact same fixed box on the
    page, whatever the new logo's native dimensions/aspect ratio are.

    Done as a post-save zip edit rather than through python-docx: the
    on-page size lives in the drawing's <wp:extent>/<a:ext> XML, which we
    never touch, so swapping only the raster bytes behind the existing
    relationship is both sufficient and the least risky way to do this.
    """
    canvas_size = _read_media_pixel_size(docx_path, media_path)
    new_bytes = compose_logo_png(logo_source_path, canvas_size)

    tmp_fd, tmp_path = tempfile.mkstemp(suffix=".docx", dir=str(Path(docx_path).parent))
    os.close(tmp_fd)
    try:
        with zipfile.ZipFile(docx_path, "r") as zin, \
                zipfile.ZipFile(tmp_path, "w", zipfile.ZIP_DEFLATED) as zout:
            for item in zin.infolist():
                data = new_bytes if item.filename == media_path else zin.read(item.filename)
                zout.writestr(item, data)
        shutil.move(tmp_path, docx_path)
    except Exception:
        if os.path.exists(tmp_path):
            os.remove(tmp_path)
        raise


# ---------------------------------------------------------------------------
# the actual fill
# ---------------------------------------------------------------------------

def build_statement(template_path: str, data: StatementData, output_path: str) -> None:
    compute_totals(data)
    doc = docx.Document(template_path)
    tables = doc.tables

    # --- table 0: statement number / date -----------------------------
    header_cell = tables[0].rows[0].cells[1]
    replace_in_cell(header_cell, {
        "[000001]": data.statement_number or "000001",
        "[MM/DD/YYYY]": data.statement_date or "",
    })

    # --- table 1: your company (left) / bill to (right) ---------------
    company_cell = tables[1].rows[0].cells[0]
    set_paragraph_text(company_cell.paragraphs[0], data.company_name or "COASTAL HPC")
    phone_email = "  |  ".join(v for v in (data.company_phone, data.company_email) if v)
    replace_in_cell(company_cell, {
        "[City, State ZIP]": data.company_address or "",
        "[Phone]  |  [Email]": phone_email,
        "[Website]": data.company_website or "",
    })

    customer_cell = tables[1].rows[0].cells[1]
    email_phone = " / ".join(v for v in (data.customer_email, data.customer_phone) if v)
    replace_in_cell(customer_cell, {
        "[Customer / Company Name]": data.customer_name or "",
        "[Contact Name]": data.contact_name or "",
        "[Billing Address]": data.billing_address or "",
        "[City, State ZIP]": data.customer_city_state_zip or "",
        "[Email / Phone]": email_phone,
    })

    # --- table 2: account summary --------------------------------------
    summary_row = tables[2].rows[1]
    for cell, value in zip(
        summary_row.cells,
        (data.previous_balance, data.payments_credits, data.new_charges,
         data.past_due, data.amount_due),
    ):
        set_cell_text(cell, money(value))

    # --- table 3: statement activity -----------------------------------
    activity_table = tables[3]
    total_row_index = len(activity_table.rows) - 1
    body_row_count = total_row_index - 1  # rows between header (0) and total row
    n_items = len(data.activity)

    if n_items > body_row_count:
        for _ in range(n_items - body_row_count):
            clone_row_before(activity_table, 1, total_row_index)
            total_row_index += 1
    elif n_items < body_row_count:
        for _ in range(body_row_count - n_items):
            remove_row(activity_table, total_row_index - 1)
            total_row_index -= 1

    for i, item in enumerate(data.activity):
        row = activity_table.rows[1 + i]
        set_cell_text(row.cells[0], item.date)
        set_cell_text(row.cells[1], item.reference)
        set_cell_text(row.cells[2], item.description)
        set_cell_text(row.cells[3], money(item.charge) if item.charge is not None else "")
        set_cell_text(row.cells[4], money(item.payment) if item.payment is not None else "")

    total_row = activity_table.rows[total_row_index]
    total_charges = sum(a.charge or 0.0 for a in data.activity)
    total_payments = sum(a.payment or 0.0 for a in data.activity)
    set_cell_text(total_row.cells[3], money(total_charges))
    set_cell_text(total_row.cells[4], money(total_payments))

    # --- table 4: payment information ----------------------------------
    payment_info_cell = tables[4].rows[0].cells[0]
    replace_in_cell(payment_info_cell, {
        "[ACH / Check / Card / Online]": data.payment_methods or "",
        "Coastal HPC": data.payable_to or "Coastal HPC",
        "[Insert payment instructions]": data.payment_instructions or "",
    })

    payment_due_cell = tables[4].rows[0].cells[1]
    replace_in_cell(payment_due_cell, {
        "[11/08/1977]": data.payment_due_date or "",
    })
    for p in payment_due_cell.paragraphs:
        for r in p.runs:
            if "$0.00" in r.text:
                r.text = r.text.replace("$0.00", money(data.amount_due))

    # --- notes / terms (a plain paragraph, not a table cell) -----------
    placeholder = ("[Add payment terms, late-fee policy, project notes, "
                    "purchase order references, or a thank-you message here.]")
    for p in doc.paragraphs:
        if placeholder in p.text:
            set_paragraph_text(p, data.notes_terms or "")
            break

    doc.save(output_path)

    if data.company_logo_path:
        replace_logo(output_path, data.company_logo_path)
