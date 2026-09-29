"""The 5 selectable color themes and 5 selectable fonts for the billing
statement, and the logic that applies them to a python-docx Document.

Inspected directly from the template's XML (see comments below for the
exact hex values found) rather than guessed: the whole design uses only
two brand colors --

  0B1E33  dark navy  -- header/section-title text, and the fill behind
                        every dark table-header row and the PAYMENT DUE box
  00BFEF  cyan        -- "BILL TO"/"PAYMENT INFORMATION" labels, the
                        AMOUNT DUE header cell's fill, and the payment due
                        date text
  EAF8FC  (a ~87%-lightened tint of 00BFEF) -- the AMOUNT DUE value cell's
                        highlight fill

-- plus a neutral dark slate (172B3A, the document's default body-text
color) and a few near-white neutral cell backgrounds (F4F7F9/F5F7F9/F7FAFC)
that are deliberately left alone by every theme: keeping ordinary body
text and plain row backgrounds neutral regardless of brand color is what
keeps the recolored document looking like a real business form rather
than everything being tinted.
"""

from __future__ import annotations

from typing import Dict, Tuple

from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import RGBColor

PRIMARY_HEX = "0B1E33"
ACCENT_HEX = "00BFEF"
ACCENT_TINT_HEX = "EAF8FC"

# name -> (primary/dark hex, accent hex). Index 0 is the template's own
# original look, kept byte-for-byte identical as the default choice.
COLOR_THEMES: Dict[str, Tuple[str, str]] = {
    "Coastal Blue (default)": ("0B1E33", "00BFEF"),
    "Forest Green": ("0B3323", "16A34A"),
    "Burgundy Red": ("331016", "DC2645"),
    "Slate Purple": ("201B33", "7C5CFF"),
    "Sunset Orange": ("331C0B", "F2760F"),
}

# name -> font family name. All five are pre-installed on every Windows
# machine, so a generated .docx never falls back to a substitute font.
FONT_CHOICES: Dict[str, str] = {
    "Arial (default)": "Arial",
    "Calibri": "Calibri",
    "Georgia": "Georgia",
    "Times New Roman": "Times New Roman",
    "Verdana": "Verdana",
}

DEFAULT_COLOR_THEME = "Coastal Blue (default)"
DEFAULT_FONT = "Arial (default)"


def lighten(hex_color: str, amount: float) -> str:
    """Blends `hex_color` toward white by `amount` (0 = unchanged, 1 =
    white). Used to derive the AMOUNT DUE highlight tint from whichever
    accent color a theme picks, instead of hand-picking a 6th color per
    theme that has to be kept visually consistent by hand."""
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (0, 2, 4))
    r = round(r + (255 - r) * amount)
    g = round(g + (255 - g) * amount)
    b = round(b + (255 - b) * amount)
    return f"{r:02X}{g:02X}{b:02X}"


def _iter_all_paragraphs(doc):
    """Every paragraph in the document worth restyling: body, every table
    cell (recursively, in case a cell ever contains a nested table), and
    every section's footer. python-docx has no single built-in method for
    this, so it's assembled once here rather than re-derived per caller."""
    for p in doc.paragraphs:
        yield p
    for t in doc.tables:
        yield from _iter_table_paragraphs(t)
    for section in doc.sections:
        for p in section.footer.paragraphs:
            yield p
        for t in section.footer.tables:
            yield from _iter_table_paragraphs(t)


def _iter_table_paragraphs(table):
    for row in table.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                yield p
            for nested in cell.tables:
                yield from _iter_table_paragraphs(nested)


def apply_font(doc, font_name: str) -> None:
    """Forces every run in the document to `font_name`. Deliberately
    touches every run directly rather than editing styles.xml's defaults:
    a handful of runs already carry an explicit font override (the
    template's own Arial/Play choices), and a direct run-level override
    always wins over a style default -- so setting it everywhere is the
    only way that's guaranteed to actually change what's on the page."""
    for p in _iter_all_paragraphs(doc):
        for run in p.runs:
            run.font.name = font_name
            rPr = run._element.get_or_add_rPr()
            rFonts = rPr.find(qn("w:rFonts"))
            if rFonts is None:
                rFonts = OxmlElement("w:rFonts")
                rPr.append(rFonts)
            rFonts.set(qn("w:eastAsia"), font_name)


def _recolor_run(run, color_map: Dict[str, str]) -> None:
    color = run.font.color
    if color is None or color.rgb is None:
        return
    current = str(color.rgb).upper()
    if current in color_map:
        run.font.color.rgb = RGBColor.from_string(color_map[current])


def _recolor_cell_fill(cell, color_map: Dict[str, str]) -> None:
    tcPr = cell._tc.tcPr
    if tcPr is None:
        return
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        return
    current = (shd.get(qn("w:fill")) or "").upper()
    if current in color_map:
        shd.set(qn("w:fill"), color_map[current])


def apply_color_theme(doc, primary_hex: str, accent_hex: str) -> None:
    """Remaps the template's two brand colors (and the accent's derived
    highlight tint) to a new theme, wherever they appear as either run
    text color or table-cell shading. Neutral text/background colors
    (body text, plain row backgrounds) are untouched by design -- see the
    module docstring."""
    # When the accent is unchanged (the default theme), reuse the
    # template's own exact tint rather than a re-derived approximation of
    # it -- otherwise the "default" theme wouldn't quite reproduce the
    # original document byte-for-byte.
    if accent_hex.upper() == ACCENT_HEX.upper():
        accent_tint = ACCENT_TINT_HEX
    else:
        accent_tint = lighten(accent_hex, 0.87)

    color_map = {
        PRIMARY_HEX.upper(): primary_hex.upper(),
        ACCENT_HEX.upper(): accent_hex.upper(),
        ACCENT_TINT_HEX.upper(): accent_tint.upper(),
    }

    for p in _iter_all_paragraphs(doc):
        for run in p.runs:
            _recolor_run(run, color_map)

    for t in doc.tables:
        _recolor_table_cells(t, color_map)
    for section in doc.sections:
        for t in section.footer.tables:
            _recolor_table_cells(t, color_map)


def _recolor_table_cells(table, color_map: Dict[str, str]) -> None:
    for row in table.rows:
        for cell in row.cells:
            _recolor_cell_fill(cell, color_map)
            for nested in cell.tables:
                _recolor_table_cells(nested, color_map)
